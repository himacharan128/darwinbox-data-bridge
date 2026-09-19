from __future__ import annotations

import csv
import pathlib

import pytest
import yaml
from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parents[2]
FIX = ROOT / "tests" / "fixtures"
CORPUS = ROOT / "tests" / "evaluations" / "corpus"


def _csv(rel: str) -> list[dict]:
    with (FIX / rel).open(encoding="utf-8") as f:
        return list(csv.DictReader(f))


@pytest.fixture(scope="session")
def schema() -> dict:
    return yaml.safe_load((FIX / "schemas" / "target_schema.yaml").read_text())


@pytest.fixture(scope="session")
def fields(schema) -> dict:
    return {f["name"]: f for f in schema["fields"]}


@pytest.fixture(scope="session")
def corpus() -> dict:
    return {p.stem: yaml.safe_load(p.read_text()) for p in CORPUS.glob("*.yaml")}


@pytest.fixture(scope="session")
def hrms() -> list[dict]:
    return _csv("run1/hrms_employees_export.csv")


@pytest.fixture(scope="session")
def contractors() -> list[dict]:
    return _csv("run1/contractors_2024.csv")


@pytest.fixture(scope="session")
def payroll() -> list[dict]:
    ws = load_workbook(FIX / "run1" / "payroll_staff.xlsx")["Staff"]
    header = [c.value for c in ws[1]]
    return [dict(zip(header, [c.value for c in row])) for row in ws.iter_rows(min_row=2)]


@pytest.fixture(scope="session")
def departments() -> list[dict]:
    return _csv("run1/departments.csv")


@pytest.fixture(scope="session")
def run2() -> list[dict]:
    return _csv("run2/hrms_delta_export.csv")


@pytest.fixture(scope="session")
def by_id(hrms) -> dict:
    return {r["Emp ID"]: r for r in hrms}
